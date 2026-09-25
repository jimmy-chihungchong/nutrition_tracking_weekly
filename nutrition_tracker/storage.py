"""Meal log stored in a plain Excel workbook (data/nutrition_log.xlsx).

Sheet "Meals": one row per meal, with fixed info columns followed by one column per
nutrient, headed like "Protein (g)". Sheet "Column Keys" records which nutrient each
column holds, so renaming a nutrient in the targets file keeps its history.

Safety:
- one lock per workbook for the whole process, so simultaneous saves can't interleave;
- saves go to a temporary file that then replaces the workbook, so a failed save never corrupts it;
- writes are refused while the workbook is open in Excel/LibreOffice, whose own save
  would otherwise overwrite meals logged in the app.
"""
import os
import tempfile
import threading
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .nutrients import Nutrient

MEALS_SHEET = "Meals"
KEYS_SHEET = "Column Keys"
INFO_COLUMNS = ["Meal ID", "Date", "Meal Type", "Meal Description", "Food Items & Quantities",
                "Estimated By", "Logged At"]
MEAL_TYPES = ["Breakfast", "Lunch", "Dinner", "Snack"]
SAMPLE_MARKER = "(sample data)"
EXCEL_CELL_LIMIT = 32000

ABOUT_TEXT = [
    "Weekly Nutrition Tracker — meal log",
    "",
    "Each row on the 'Meals' sheet is one logged meal.",
    "Date / Meal Type / Meal Description: what the user entered.",
    "Food Items & Quantities: how the meal was interpreted (item, quantity, estimated grams).",
    "Estimated By: which nutrition engine produced the numbers. '(sample data)' marks demo meals.",
    "Remaining columns: estimated amount of each nutrient in that meal, unit shown in the header.",
    "'Column Keys' sheet: which nutrient each column holds (used by the app — please don't edit).",
    "",
    "Weekly totals are the sum of these columns for Monday–Sunday.",
    "Goals and targets are kept separately in data/weekly_targets.csv.",
    "Close this file before logging meals in the app, otherwise the app can't save.",
    "Estimates are approximate and not clinically validated.",
]

_locks: dict[str, threading.RLock] = {}
_locks_guard = threading.Lock()


def _lock_for(path: Path) -> threading.RLock:
    with _locks_guard:
        return _locks.setdefault(str(path.resolve()).lower(), threading.RLock())


class StorageError(Exception):
    pass


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 20000 < value < 80000:
        return date(1899, 12, 30) + timedelta(days=int(value))  # spreadsheet date serial number
    if isinstance(value, str) and value.strip():
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


def _clip(text: str) -> str:
    return text if len(text) <= EXCEL_CELL_LIMIT else text[:EXCEL_CELL_LIMIT] + "…"


def nutrient_columns(header: list, keymap: dict[str, str], nutrients: list[Nutrient]) -> dict[str, int]:
    """nutrient key -> 0-based column index, matched by key (via the Column Keys sheet) first, then by label."""
    cols = {}
    for n in nutrients:
        h = keymap.get(n.key)
        if h in header:
            cols[n.key] = header.index(h)
        elif n.label in header:
            cols[n.key] = header.index(n.label)
    return cols


def parse_meal_rows(rows: list, keymap: dict[str, str], nutrients: list[Nutrient],
                    start: date | None = None, end: date | None = None) -> tuple[list[dict], bool]:
    """Turn a meal table (header row + data rows) into meal dicts. Shared by the Excel and Google Sheets stores.

    Returns (meals sorted by date and meal type, whether any dated row is missing a Meal ID).
    """
    if not rows:
        return [], False
    header = list(rows[0])
    col = {h: i for i, h in enumerate(header) if h}
    nut_cols = nutrient_columns(header, keymap, nutrients)
    get = lambda r, h: r[col[h]] if h in col and col[h] < len(r) else None

    meals, missing_ids = [], False
    for r in rows[1:]:
        d = _as_date(get(r, "Date"))
        if d is None:
            continue
        if not get(r, "Meal ID"):
            missing_ids = True  # shown without an ID (not editable) until the ID can be backfilled
        if (start and d < start) or (end and d > end):
            continue
        values = {}
        for n in nutrients:
            i = nut_cols.get(n.key)
            try:
                values[n.key] = float(r[i] or 0) if i is not None and i < len(r) else 0.0
            except (TypeError, ValueError):
                values[n.key] = 0.0
        logged = get(r, "Logged At")
        meals.append({
            "id": str(get(r, "Meal ID") or ""),
            "date": d.isoformat(),
            "meal_type": str(get(r, "Meal Type") or ""),
            "description": str(get(r, "Meal Description") or ""),
            "items_text": str(get(r, "Food Items & Quantities") or ""),
            "source": str(get(r, "Estimated By") or ""),
            "logged_at": logged.isoformat() if isinstance(logged, datetime) else str(logged or ""),
            "nutrients": values,
        })
    order = {t: i for i, t in enumerate(MEAL_TYPES)}
    meals.sort(key=lambda m: (m["date"], order.get(m["meal_type"], 9), m["logged_at"]))
    return meals, missing_ids


class ExcelMealStore:
    def __init__(self, path, nutrients: list[Nutrient]):
        self.path = Path(path)
        self.nutrients = nutrients
        self._lock = _lock_for(self.path)

    @property
    def location(self) -> str:
        return f"data/{self.path.name}"

    # ------------------------------------------------------------------ file handling
    def open_elsewhere(self) -> Path | None:
        """Return the lock file if the workbook is open in LibreOffice (.~lock.x#) or Excel (~$x)."""
        for name in (f".~lock.{self.path.name}#", f"~${self.path.name}"):
            candidate = self.path.parent / name
            if candidate.exists():
                return candidate
        return None

    def _check_writable(self) -> None:
        lock = self.open_elsewhere()
        if lock:
            raise StorageError(
                f"{self.path.name} is open in a spreadsheet program. Close it first so your changes "
                f"aren't overwritten. (If it isn't open, delete the leftover file '{lock.name}' in the data folder.)")

    def _save(self, wb) -> None:
        fd, tmp = tempfile.mkstemp(dir=self.path.parent, suffix=".xlsx.tmp")
        os.close(fd)
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError as e:
            raise StorageError(
                f"Couldn't write to {self.path.name} — it is probably open in Excel. Close it and try again.") from e
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)

    def _load(self, **kwargs):
        try:
            return load_workbook(self.path, **kwargs)
        except PermissionError as e:
            raise StorageError(f"Couldn't open {self.path.name}. Close it in Excel and try again.") from e
        except (BadZipFile, KeyError, OSError) as e:
            raise StorageError(f"{self.path.name} couldn't be read ({e}). Restore it from a backup.") from e

    def ensure_workbook(self) -> bool:
        """Create the workbook with headers if it doesn't exist. Returns True if created."""
        with self._lock:
            if self.path.exists():
                return False
            self.path.parent.mkdir(parents=True, exist_ok=True)
            wb = Workbook()
            ws = wb.active
            ws.title = MEALS_SHEET
            ws.append(INFO_COLUMNS + [n.label for n in self.nutrients])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2F6B4F")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 32
            ws.freeze_panes = "E2"
            for col, width in zip("ABCDEFG", (11, 12, 11, 48, 48, 30, 19)):
                ws.column_dimensions[col].width = width
            about = wb.create_sheet("About")
            for line in ABOUT_TEXT:
                about.append([line])
            about["A1"].font = Font(bold=True, size=13)
            about.column_dimensions["A"].width = 100
            self._write_keymap(wb, {n.key: n.label for n in self.nutrients})
            self._save(wb)
            return True

    # ------------------------------------------------------------------ column mapping
    @staticmethod
    def _read_keymap(wb) -> dict[str, str]:
        if KEYS_SHEET not in wb.sheetnames:
            return {}
        mapping = {}
        for row in list(wb[KEYS_SHEET].iter_rows(values_only=True))[1:]:
            if row and len(row) >= 2 and row[0] and row[1]:
                mapping[str(row[0])] = str(row[1])
        return mapping

    @staticmethod
    def _write_keymap(wb, mapping: dict[str, str]) -> None:
        if KEYS_SHEET in wb.sheetnames:
            del wb[KEYS_SHEET]
        ws = wb.create_sheet(KEYS_SHEET)
        ws.append(["Nutrient key", "Column header on the Meals sheet"])
        ws["A1"].font = ws["B1"].font = Font(bold=True)
        for key, header in mapping.items():
            ws.append([key, header])
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 36

    def _nutrient_columns(self, header: list, keymap: dict[str, str]) -> dict[str, int]:
        return nutrient_columns(header, keymap, self.nutrients)

    def _prepare_for_write(self, wb):
        """Make sure every nutrient has a column (renaming headers if a nutrient was renamed)."""
        ws = wb[MEALS_SHEET]
        header = [c.value for c in ws[1]]
        keymap = self._read_keymap(wb)
        cols = self._nutrient_columns(header, keymap)
        for n in self.nutrients:
            if n.key not in cols:
                header.append(n.label)
                cols[n.key] = len(header) - 1
                ws.cell(row=1, column=len(header), value=n.label).font = Font(bold=True)
            elif header[cols[n.key]] != n.label:
                header[cols[n.key]] = n.label
                ws.cell(row=1, column=cols[n.key] + 1, value=n.label)
        keymap.update({n.key: n.label for n in self.nutrients})
        self._write_keymap(wb, keymap)
        return ws, header, cols

    # ------------------------------------------------------------------ writes
    def _fill_row(self, ws, row: int, header: list, cols: dict, meal: dict) -> None:
        info = {
            "Meal ID": meal["id"], "Date": meal["date"], "Meal Type": meal["meal_type"],
            "Meal Description": _clip(meal["description"]), "Food Items & Quantities": _clip(meal["items_text"]),
            "Estimated By": _clip(meal["source"]), "Logged At": meal["logged_at"],
        }
        for name, value in info.items():
            if name in header:
                cell = ws.cell(row=row, column=header.index(name) + 1, value=value)
                if name == "Date":
                    cell.number_format = "yyyy-mm-dd"
                elif name == "Logged At":
                    cell.number_format = "yyyy-mm-dd hh:mm"
        for n in self.nutrients:
            ws.cell(row=row, column=cols[n.key] + 1, value=round(float(meal["nutrients"].get(n.key, 0.0)), 3))

    def _new_meal(self, meal_date: date, meal_type: str, description: str, items_text: str,
                  source: str, nutrients: dict) -> dict:
        return {"id": uuid.uuid4().hex[:8], "date": meal_date, "meal_type": meal_type, "description": description,
                "items_text": items_text, "source": source, "logged_at": datetime.now().replace(microsecond=0),
                "nutrients": {n.key: float(nutrients.get(n.key, 0.0)) for n in self.nutrients}}

    @staticmethod
    def _public(meal: dict) -> dict:
        return {**meal, "date": meal["date"].isoformat(), "logged_at": meal["logged_at"].isoformat()}

    def add_meals(self, entries: list[dict]) -> list[dict]:
        """Add several meals in one save. Each entry has the keyword arguments of add_meal()."""
        meals = [self._new_meal(**e) for e in entries]
        with self._lock:
            self.ensure_workbook()
            self._check_writable()
            wb = self._load()
            ws, header, cols = self._prepare_for_write(wb)
            for meal in meals:
                self._fill_row(ws, ws.max_row + 1, header, cols, meal)
            self._save(wb)
        return [self._public(m) for m in meals]

    def add_meal(self, *, meal_date: date, meal_type: str, description: str, items_text: str,
                 source: str, nutrients: dict[str, float]) -> dict:
        return self.add_meals([dict(meal_date=meal_date, meal_type=meal_type, description=description,
                                    items_text=items_text, source=source, nutrients=nutrients)])[0]

    def _find_row(self, ws, header: list, meal_id: str) -> int | None:
        if "Meal ID" not in header:
            return None
        col = header.index("Meal ID") + 1
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=col).value) == meal_id:
                return row
        return None

    def update_meal(self, meal_id: str, *, meal_date: date, meal_type: str, description: str, items_text: str,
                    source: str, nutrients: dict[str, float]) -> dict | None:
        with self._lock:
            if not self.path.exists():
                return None
            self._check_writable()
            wb = self._load()
            ws, header, cols = self._prepare_for_write(wb)
            row = self._find_row(ws, header, meal_id)
            if not row:
                return None
            logged = ws.cell(row=row, column=header.index("Logged At") + 1).value if "Logged At" in header else None
            meal = self._new_meal(meal_date, meal_type, description, items_text, source, nutrients)
            meal["id"] = meal_id
            if isinstance(logged, datetime):
                meal["logged_at"] = logged
            self._fill_row(ws, row, header, cols, meal)
            self._save(wb)
        return self._public(meal)

    def delete_meal(self, meal_id: str) -> bool:
        with self._lock:
            if not self.path.exists():
                return False
            self._check_writable()
            wb = self._load()
            ws = wb[MEALS_SHEET]
            row = self._find_row(ws, [c.value for c in ws[1]], meal_id)
            if not row:
                return False
            ws.delete_rows(row)
            self._save(wb)
        return True

    def delete_sample_meals(self) -> int:
        """Remove only the demo meals (marked '(sample data)'); the user's own meals are kept."""
        with self._lock:
            if not self.path.exists():
                return 0
            self._check_writable()
            wb = self._load()
            ws = wb[MEALS_SHEET]
            header = [c.value for c in ws[1]]
            if "Estimated By" not in header:
                return 0
            col = header.index("Estimated By") + 1
            rows = [r for r in range(2, ws.max_row + 1)
                    if str(ws.cell(row=r, column=col).value or "").endswith(SAMPLE_MARKER)]
            for r in reversed(rows):
                ws.delete_rows(r)
            if rows:
                self._save(wb)
            return len(rows)

    def _backfill_ids(self) -> None:
        """Give rows typed into Excel by hand a proper Meal ID, so they can be edited/deleted safely."""
        with self._lock:
            if self.open_elsewhere():
                return
            wb = self._load()
            ws = wb[MEALS_SHEET]
            header = [c.value for c in ws[1]]
            if "Meal ID" not in header or "Date" not in header:
                return
            id_col, date_col = header.index("Meal ID") + 1, header.index("Date") + 1
            changed = False
            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=date_col).value and not ws.cell(row=row, column=id_col).value:
                    ws.cell(row=row, column=id_col, value=uuid.uuid4().hex[:8])
                    changed = True
            if changed:
                self._save(wb)

    # ------------------------------------------------------------------ reads
    def list_meals(self, start: date | None = None, end: date | None = None, _retry: bool = True) -> list[dict]:
        if not self.path.exists():
            return []
        with self._lock:
            wb = self._load(read_only=True, data_only=True)
            try:
                rows = list(wb[MEALS_SHEET].iter_rows(values_only=True))
                keymap = self._read_keymap(wb)
            finally:
                wb.close()
        meals, missing_ids = parse_meal_rows(rows, keymap, self.nutrients, start, end)
        if missing_ids and _retry and not self.open_elsewhere():
            self._backfill_ids()
            return self.list_meals(start, end, _retry=False)
        return meals
